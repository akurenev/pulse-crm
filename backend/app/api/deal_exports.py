"""Role-authorized, bounded monthly deal downloads without stored export files."""

from __future__ import annotations

import json
import uuid
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Literal
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import anyio
import sqlalchemy as sa
from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import get_session
from app.models import (
    Company,
    Contact,
    CustomFieldDefinition,
    Deal,
    DealContact,
    FieldEntity,
    Membership,
    Pipeline,
    Source,
    Stage,
    User,
)
from app.security import CurrentCRMExporter
from app.services.events import record_audit_event

router = APIRouter(tags=["crm"])
MAX_EXPORT_DEALS = 10_000
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class DealExportRequest(BaseModel):
    month: str = Field(pattern=r"^[0-9]{4}-(0[1-9]|1[0-2])$")
    date_field: Literal["created_at", "next_purchase_at", "updated_at"] = "created_at"
    pipeline_id: uuid.UUID | None = None


def _month_bounds(month: str, timezone: ZoneInfo) -> tuple[datetime, datetime]:
    year, number = map(int, month.split("-"))
    try:
        start = datetime(year, number, 1, tzinfo=timezone)
        end = datetime(year + (number == 12), number % 12 + 1, 1, tzinfo=timezone)
        return start.astimezone(UTC), end.astimezone(UTC)
    except (ValueError, OverflowError) as exc:
        raise HTTPException(status_code=422, detail="invalid export month") from exc


def _cell_value(value: Any, timezone: ZoneInfo) -> Any:
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or UTC).astimezone(timezone).replace(tzinfo=None)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str):
        # Excel cannot represent longer cells. Fail explicitly instead of silently
        # truncating a user's fields; control characters are invalid in OOXML.
        if len(value) > 32_767:
            raise HTTPException(status_code=422, detail={"code": "export_cell_too_long"})
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _workbook(headers: list[str], rows: list[list[Any]], timezone: ZoneInfo) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Сделки"
    sheet.freeze_panes = "C2"
    sheet.append([_cell_value(header, timezone) for header in headers])
    for cell in sheet[1]:
        cell.data_type = "s"
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 42
    for row_number, row in enumerate(rows, 2):
        sheet.append([_cell_value(value, timezone) for value in row])
        for column_number in range(1, len(headers) + 1):
            cell = sheet.cell(row_number, column_number)
            if isinstance(cell.value, str):
                # Force untrusted CRM text to string, including '=' and Excel
                # error literals. No field is ever interpreted as a formula.
                cell.data_type = "s"
            if isinstance(cell.value, datetime):
                cell.number_format = "dd.mm.yyyy hh:mm"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(48, max(20, len(header) + 2))
    sheet.auto_filter.ref = sheet.dimensions
    workbook.properties.description = f"Даты в часовом поясе {timezone.key}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.post("/deals/export")
async def export_deals(
    payload: DealExportRequest,
    context: CurrentCRMExporter,
    db: AsyncSession = Depends(get_session),
) -> Response:
    try:
        timezone = ZoneInfo(context.workspace.timezone)
    except ZoneInfoNotFoundError:
        timezone = ZoneInfo("UTC")
    start, end = _month_bounds(payload.month, timezone)
    date_column = getattr(Deal, payload.date_field)
    query = (
        sa.select(Deal, Pipeline.name, Stage.name, Company, User, Source)
        .join(
            Pipeline,
            sa.and_(Pipeline.id == Deal.pipeline_id, Pipeline.workspace_id == context.workspace_id),
        )
        .join(Stage, sa.and_(Stage.id == Deal.stage_id, Stage.workspace_id == context.workspace_id))
        .outerjoin(
            Company,
            sa.and_(
                Company.id == Deal.company_id,
                Company.workspace_id == context.workspace_id,
                Company.deleted_at.is_(None),
            ),
        )
        .outerjoin(
            User,
            sa.and_(
                User.id == Deal.assignee_id,
                User.id.in_(
                    sa.select(Membership.user_id).where(
                        Membership.workspace_id == context.workspace_id
                    )
                ),
            ),
        )
        .outerjoin(
            Source,
            sa.and_(Source.id == Deal.source_id, Source.workspace_id == context.workspace_id),
        )
        .where(
            Deal.workspace_id == context.workspace_id,
            Deal.deleted_at.is_(None),
            date_column >= start,
            date_column < end,
        )
        .order_by(date_column, Deal.id)
        .limit(MAX_EXPORT_DEALS + 1)
    )
    if payload.pipeline_id:
        query = query.where(Deal.pipeline_id == payload.pipeline_id)
    records = (await db.execute(query)).all()
    if len(records) > MAX_EXPORT_DEALS:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "export_row_limit",
                "limit": MAX_EXPORT_DEALS,
            },
        )
    deal_ids = [record[0].id for record in records]
    contacts_by_deal: dict[Any, list[Contact]] = {}
    if deal_ids:
        contact_rows = (
            await db.execute(
                sa.select(DealContact.deal_id, Contact)
                .join(Contact, Contact.id == DealContact.contact_id)
                .where(
                    DealContact.workspace_id == context.workspace_id,
                    DealContact.deal_id.in_(deal_ids),
                    Contact.workspace_id == context.workspace_id,
                    Contact.deleted_at.is_(None),
                )
                .order_by(DealContact.created_at, Contact.id)
            )
        ).all()
        for deal_id, contact in contact_rows:
            contacts_by_deal.setdefault(deal_id, []).append(contact)
    definitions = (
        await db.scalars(
            sa.select(CustomFieldDefinition)
            .where(
                CustomFieldDefinition.workspace_id == context.workspace_id,
                CustomFieldDefinition.entity_type == FieldEntity.deal,
            )
            .order_by(CustomFieldDefinition.created_at, CustomFieldDefinition.key)
        )
    ).all()
    field_names = {definition.key: definition.name for definition in definitions}
    extra_keys = sorted(
        {key for record in records for key in record[0].custom_fields} - field_names.keys()
    )
    keys = [*field_names, *extra_keys]
    headers = [
        "ID сделки",
        "Название",
        "Сумма",
        "Валюта",
        "Воронка",
        "ID воронки",
        "Этап",
        "ID этапа",
        "Компания",
        "ID компании",
        "ИНН компании",
        "Телефон компании",
        "Email компании",
        "Контакты",
        "ID контактов",
        "Телефоны контактов",
        "Email контактов",
        "Ответственный",
        "ID ответственного",
        "Email ответственного",
        "Источник",
        "ID источника",
        "Код источника",
        "Теги",
        "Следующая покупка",
        "Последняя активность",
        "Создана",
        "Изменена",
        "Версия",
        "Пользовательские поля (JSON)",
        *(f"{field_names.get(key, key)} [{key}]" for key in keys),
    ]
    rows: list[list[Any]] = []
    for deal, pipeline_name, stage_name, company, assignee, source in records:
        contacts = contacts_by_deal.get(deal.id, [])
        rows.append(
            [
                str(deal.id),
                deal.title,
                deal.amount,
                deal.currency,
                pipeline_name,
                str(deal.pipeline_id),
                stage_name,
                str(deal.stage_id),
                company.name if company else None,
                str(deal.company_id) if deal.company_id else None,
                company.inn if company else None,
                company.phone if company else None,
                company.email if company else None,
                "\n".join(
                    f"{contact.first_name} {contact.last_name}".strip() for contact in contacts
                ),
                "\n".join(str(contact.id) for contact in contacts),
                "\n".join(
                    "; ".join(
                        dict.fromkeys(
                            value for value in [contact.primary_phone, *contact.phones] if value
                        )
                    )
                    for contact in contacts
                ),
                "\n".join(
                    "; ".join(
                        dict.fromkeys(
                            value for value in [contact.primary_email, *contact.emails] if value
                        )
                    )
                    for contact in contacts
                ),
                assignee.full_name if assignee else None,
                str(deal.assignee_id) if deal.assignee_id else None,
                assignee.email if assignee else None,
                source.name if source else None,
                str(deal.source_id) if deal.source_id else None,
                source.key if source else None,
                deal.tags,
                deal.next_purchase_at,
                deal.last_activity_at,
                deal.created_at,
                deal.updated_at,
                deal.version,
                deal.custom_fields,
                *(deal.custom_fields.get(key) for key in keys),
            ]
        )
    content = await anyio.to_thread.run_sync(_workbook, headers, rows, timezone)
    record_audit_event(
        db,
        workspace_id=context.workspace_id,
        event_type="deals.export.downloaded",
        entity_type="workspace",
        entity_id=context.workspace_id,
        actor_id=context.user_id,
        payload={
            "month": payload.month,
            "date_field": payload.date_field,
            "pipeline_id": str(payload.pipeline_id) if payload.pipeline_id else None,
            "count": len(records),
        },
    )
    await db.commit()
    return Response(
        content,
        media_type=XLSX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="deals-{payload.month}.xlsx"',
            "Cache-Control": "no-store",
        },
    )
