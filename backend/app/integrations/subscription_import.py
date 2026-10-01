"""Import expiring subscription spreadsheets into grouped renewal deals."""

from __future__ import annotations

import hashlib
import re
import uuid
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time, timedelta, tzinfo
from io import BytesIO
from typing import Any
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import sqlalchemy as sa
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]
from sqlalchemy.ext.asyncio import AsyncSession

from app.integrations.identity import (
    IdentityNormalizationError,
    normalize_email_address,
    normalize_phone_number,
    sync_contact_points,
)
from app.integrations.models import ContactPoint, ContactPointKind, ExternalEntityMap
from app.models import (
    Company,
    Contact,
    Deal,
    DealContact,
    DealStageHistory,
    Source,
    Stage,
    StageType,
    Task,
    TaskStatus,
)
from app.services.events import record_domain_event

PROVIDER = "frontol_endlic"
FORMAT_KEY = "frontol_endlic_v1"
SOURCE_KEY = "subscription_import"
SOURCE_NAME = "Импорт подписок"
MAX_ROWS = 10_000
MAX_COLUMNS = 50
MAX_ARCHIVE_FILES = 1_000
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024

EXPECTED_HEADERS = (
    "LID",
    "Дата активации",
    "Дата окончания",
    "Наименование пользователя",
    "ИНН Пользователя",
    "Телефон пользователя",
    "Электронная почта пользователя",
)


class SubscriptionImportError(ValueError):
    """A user-correctable spreadsheet or routing error."""


@dataclass(frozen=True, slots=True)
class SubscriptionLicense:
    row_number: int
    lid: str
    activated_at: date
    expires_at: date
    customer_name: str
    inn: str
    phone: str
    normalized_phone: str
    email: str
    normalized_email: str
    product: str

    def public_dict(self) -> dict[str, Any]:
        return {
            "lid": self.lid,
            "product": self.product,
            "activated_at": self.activated_at.isoformat(),
            "expires_at": self.expires_at.isoformat(),
        }


@dataclass(frozen=True, slots=True)
class ParsedSubscriptions:
    licenses: tuple[SubscriptionLicense, ...]
    duplicate_rows: int
    warnings: tuple[str, ...]
    file_sha256: str


@dataclass(slots=True)
class ContactCluster:
    licenses: list[SubscriptionLicense] = field(default_factory=list)
    normalized_phones: set[str] = field(default_factory=set)
    normalized_emails: set[str] = field(default_factory=set)

    def overlaps(self, license: SubscriptionLicense) -> bool:
        return bool(
            (license.normalized_phone and license.normalized_phone in self.normalized_phones)
            or (license.normalized_email and license.normalized_email in self.normalized_emails)
        )

    def add(self, license: SubscriptionLicense) -> None:
        self.licenses.append(license)
        if license.normalized_phone:
            self.normalized_phones.add(license.normalized_phone)
        if license.normalized_email:
            self.normalized_emails.add(license.normalized_email)

    def merge(self, other: ContactCluster) -> None:
        for license in other.licenses:
            self.add(license)

    @property
    def name(self) -> str:
        names = [item.customer_name for item in self.licenses if item.customer_name]
        return max(names, key=len) if names else "Контакт не указан"

    @property
    def phones(self) -> list[str]:
        return _ordered_unique(item.phone for item in self.licenses if item.phone)

    @property
    def emails(self) -> list[str]:
        return _ordered_unique(item.email for item in self.licenses if item.email)


@dataclass(frozen=True, slots=True)
class RenewalGroup:
    inn: str
    expires_at: date
    organization_name: str
    licenses: tuple[SubscriptionLicense, ...]

    @property
    def external_id(self) -> str:
        return f"{self.inn}:{self.expires_at.isoformat()}"

    @property
    def products(self) -> list[str]:
        return _ordered_unique(item.product for item in self.licenses if item.product)

    @property
    def fingerprint(self) -> str:
        payload = "\n".join(
            f"{item.lid}|{item.product}|{item.activated_at.isoformat()}|{item.expires_at.isoformat()}"
            for item in sorted(self.licenses, key=lambda value: value.lid)
        )
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()


@dataclass(frozen=True, slots=True)
class ImportGroupResult:
    organization_name: str
    inn: str
    expires_at: date
    license_count: int
    products: tuple[str, ...]
    contact_count: int
    action: str
    error: str | None = None


@dataclass(frozen=True, slots=True)
class SubscriptionImportOutcome:
    counts: dict[str, int]
    groups: tuple[ImportGroupResult, ...]
    warnings: tuple[str, ...]


def parse_frontol_endlic(content: bytes) -> ParsedSubscriptions:
    if not content.startswith(b"PK\x03\x04"):
        raise SubscriptionImportError("Файл не является корректным XLSX.")
    _validate_xlsx_container(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise SubscriptionImportError("Не удалось прочитать XLSX-файл.") from exc
    try:
        if not workbook.worksheets:
            raise SubscriptionImportError("В файле нет листов.")
        worksheet = workbook.worksheets[0]
        if worksheet.max_row > MAX_ROWS or worksheet.max_column > MAX_COLUMNS:
            raise SubscriptionImportError(
                f"Таблица превышает ограничение {MAX_ROWS} строк и {MAX_COLUMNS} столбцов."
            )
        header_row = _find_header_row(worksheet)
        current_product = ""
        licenses_by_key: dict[tuple[str, date], SubscriptionLicense] = {}
        duplicate_rows = 0
        warnings: list[str] = []
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            header_row + 1,
        ):
            values = list(row[:7]) + [None] * max(0, 7 - len(row))
            nonempty = [_cell_text(value) for value in values if _cell_text(value)]
            if len(nonempty) == 1:
                current_product = nonempty[0]
                continue
            activated_at = _parse_date(values[1])
            expires_at = _parse_date(values[2])
            if activated_at is None and expires_at is None:
                continue
            if activated_at is None or expires_at is None:
                warnings.append(f"Строка {row_number}: некорректная дата активации или окончания.")
                continue
            if any(_is_formula(value) for value in values):
                warnings.append(
                    f"Строка {row_number}: формулы в импортируемых полях не поддерживаются."
                )
                continue
            try:
                license = _parse_license_row(
                    row_number=row_number,
                    values=values,
                    activated_at=activated_at,
                    expires_at=expires_at,
                    product=current_product,
                )
            except SubscriptionImportError as exc:
                warnings.append(f"Строка {row_number}: {exc}")
                continue
            license_key = (license.lid, license.expires_at)
            existing = licenses_by_key.get(license_key)
            if existing is None:
                licenses_by_key[license_key] = license
            elif _license_identity(existing) == _license_identity(license):
                duplicate_rows += 1
            else:
                warnings.append(
                    f"Строка {row_number}: LID с этой датой окончания уже встречался "
                    "с другими данными и пропущен."
                )
        if not licenses_by_key:
            raise SubscriptionImportError("В таблице не найдено корректных строк с лицензиями.")
        return ParsedSubscriptions(
            licenses=tuple(licenses_by_key.values()),
            duplicate_rows=duplicate_rows,
            warnings=tuple(warnings),
            file_sha256=hashlib.sha256(content).hexdigest(),
        )
    finally:
        workbook.close()


def renewal_groups(licenses: tuple[SubscriptionLicense, ...]) -> tuple[RenewalGroup, ...]:
    grouped: dict[tuple[str, date], list[SubscriptionLicense]] = defaultdict(list)
    for license in licenses:
        grouped[(license.inn, license.expires_at)].append(license)
    result: list[RenewalGroup] = []
    for (inn, expires_at), items in sorted(grouped.items(), key=lambda value: value[0]):
        names = [item.customer_name for item in items if item.customer_name]
        organization_name = max(names, key=len) if names else f"Клиент с ИНН {inn}"
        result.append(
            RenewalGroup(
                inn=inn,
                expires_at=expires_at,
                organization_name=organization_name,
                licenses=tuple(sorted(items, key=lambda value: value.lid)),
            )
        )
    return tuple(result)


def contact_clusters(licenses: list[SubscriptionLicense]) -> list[ContactCluster]:
    clusters: list[ContactCluster] = []
    for license in licenses:
        matches = [cluster for cluster in clusters if cluster.overlaps(license)]
        if not matches:
            cluster = ContactCluster()
            cluster.add(license)
            clusters.append(cluster)
            continue
        target = matches[0]
        target.add(license)
        for duplicate in matches[1:]:
            target.merge(duplicate)
            clusters.remove(duplicate)
    return clusters


def default_call_due_at(*, timezone_name: str, now: datetime | None = None) -> datetime:
    try:
        target_timezone: tzinfo = ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError:
        target_timezone = UTC
    local_now = (now or datetime.now(UTC)).astimezone(target_timezone)
    if local_now.weekday() < 5 and local_now.time() < time(hour=17):
        local_due = datetime.combine(local_now.date(), time(hour=17), target_timezone)
    else:
        next_day = local_now.date() + timedelta(days=1)
        while next_day.weekday() >= 5:
            next_day += timedelta(days=1)
        local_due = datetime.combine(next_day, time(hour=10), target_timezone)
    return local_due.astimezone(UTC)


async def import_subscriptions(
    session: AsyncSession,
    *,
    workspace_id: uuid.UUID,
    actor_id: uuid.UUID,
    pipeline_id: uuid.UUID,
    stage_id: uuid.UUID,
    assignee_id: uuid.UUID,
    call_due_at: datetime,
    parsed: ParsedSubscriptions,
    dry_run: bool,
) -> SubscriptionImportOutcome:
    groups = renewal_groups(parsed.licenses)
    by_inn: dict[str, list[SubscriptionLicense]] = defaultdict(list)
    for license in parsed.licenses:
        by_inn[license.inn].append(license)

    counts: dict[str, int] = {
        "rows": len(parsed.licenses) + parsed.duplicate_rows,
        "licenses": len(parsed.licenses),
        "duplicate_rows": parsed.duplicate_rows,
        "groups": len(groups),
        "companies_created": 0,
        "companies_reused": 0,
        "contacts_created": 0,
        "contacts_reused": 0,
        "deals_created": 0,
        "deals_updated": 0,
        "deals_skipped": 0,
        "tasks_created": 0,
        "tasks_reused": 0,
        "tasks_skipped": 0,
        "errors": 0,
    }
    company_by_inn: dict[str, Company | None] = {}
    company_errors: dict[str, str] = {}
    contacts_by_license: dict[tuple[str, str, date], Contact | None] = {}
    contact_count_by_inn: dict[str, int] = {}

    for inn, licenses in by_inn.items():
        companies = list(
            (
                await session.scalars(
                    sa.select(Company).where(
                        Company.workspace_id == workspace_id,
                        Company.deleted_at.is_(None),
                        Company.inn == inn,
                    )
                )
            ).all()
        )
        if len(companies) > 1:
            company_errors[inn] = "В CRM найдено несколько организаций с этим ИНН."
            company_by_inn[inn] = None
            continue
        if companies:
            company = companies[0]
            company_by_inn[inn] = company
            counts["companies_reused"] += 1
        else:
            counts["companies_created"] += 1
            if dry_run:
                company_by_inn[inn] = None
                company = None
            else:
                name = max(
                    (item.customer_name for item in licenses if item.customer_name),
                    key=len,
                    default=f"Клиент с ИНН {inn}",
                )
                company = Company(
                    workspace_id=workspace_id,
                    name=name,
                    inn=inn,
                    phone=next((item.phone for item in licenses if item.phone), None),
                    email=next((item.email for item in licenses if item.email), None),
                    tags=["Импорт подписок", "Frontol"],
                )
                session.add(company)
                await session.flush()
                company_by_inn[inn] = company
                record_domain_event(
                    session,
                    workspace_id=workspace_id,
                    event_type="company.created",
                    entity_type="company",
                    entity_id=company.id,
                    actor_id=actor_id,
                )

        clusters = contact_clusters(licenses)
        contact_count_by_inn[inn] = len(clusters)
        for cluster in clusters:
            contact: Contact | None = None
            if company is not None:
                matches = await _matching_company_contacts(
                    session,
                    workspace_id=workspace_id,
                    company_id=company.id,
                    cluster=cluster,
                )
                if len(matches) > 1:
                    company_errors[inn] = (
                        "Контактные данные соответствуют нескольким контактам организации."
                    )
                    continue
                if matches:
                    contact = matches[0]
                    counts["contacts_reused"] += 1
                    if not dry_run:
                        await _merge_contact_points(session, contact, cluster, assignee_id)
            if contact is None and inn not in company_errors:
                counts["contacts_created"] += 1
                if not dry_run and company is not None:
                    first_name, last_name = _split_contact_name(cluster.name)
                    contact = Contact(
                        workspace_id=workspace_id,
                        company_id=company.id,
                        assignee_id=assignee_id,
                        first_name=first_name,
                        last_name=last_name,
                        primary_email=cluster.emails[0] if cluster.emails else None,
                        primary_phone=cluster.phones[0] if cluster.phones else None,
                        emails=cluster.emails,
                        phones=cluster.phones,
                        tags=["Импорт подписок", "Frontol"],
                    )
                    session.add(contact)
                    await session.flush()
                    await sync_contact_points(session, contact)
                    record_domain_event(
                        session,
                        workspace_id=workspace_id,
                        event_type="contact.created",
                        entity_type="contact",
                        entity_id=contact.id,
                        actor_id=actor_id,
                    )
            for license in cluster.licenses:
                contacts_by_license[(inn, license.lid, license.expires_at)] = contact

    source: Source | None = None
    if not dry_run:
        source = await session.scalar(
            sa.select(Source).where(
                Source.workspace_id == workspace_id,
                Source.key == SOURCE_KEY,
            )
        )
        if source is None:
            source = Source(
                workspace_id=workspace_id,
                key=SOURCE_KEY,
                name=SOURCE_NAME,
                is_active=True,
            )
            session.add(source)
            await session.flush()

    results: list[ImportGroupResult] = []
    for group in groups:
        error = company_errors.get(group.inn)
        if error:
            counts["errors"] += 1
            results.append(
                _group_result(group, contact_count_by_inn.get(group.inn, 0), "error", error)
            )
            continue
        company = company_by_inn[group.inn]
        group_contacts = _unique_contacts(
            contacts_by_license.get((group.inn, license.lid, license.expires_at))
            for license in group.licenses
        )
        existing_deal = await _mapped_entity(
            session,
            workspace_id=workspace_id,
            entity_type="renewal_group",
            external_id=group.external_id,
            model=Deal,
        )
        existing_stage = (
            await session.get(Stage, existing_deal.stage_id) if existing_deal is not None else None
        )
        finalized = existing_stage is not None and existing_stage.stage_type is not StageType.open
        if existing_deal is None:
            action = "create"
            counts["deals_created"] += 1
        elif finalized:
            action = "skip"
            counts["deals_skipped"] += 1
        else:
            action = "update"
            counts["deals_updated"] += 1

        task = await _mapped_entity(
            session,
            workspace_id=workspace_id,
            entity_type="renewal_call",
            external_id=group.external_id,
            model=Task,
        )
        if finalized:
            if task is None:
                counts["tasks_skipped"] += 1
            else:
                counts["tasks_reused"] += 1
        elif task is None:
            counts["tasks_created"] += 1
        else:
            counts["tasks_reused"] += 1

        if dry_run:
            results.append(
                _group_result(group, contact_count_by_inn.get(group.inn, 0), action, None)
            )
            continue
        if finalized:
            results.append(
                _group_result(group, contact_count_by_inn.get(group.inn, 0), action, None)
            )
            continue
        if company is None or source is None:  # pragma: no cover - guarded above
            raise RuntimeError("subscription import references were not created")

        if existing_deal is None:
            deal = Deal(
                workspace_id=workspace_id,
                pipeline_id=pipeline_id,
                stage_id=stage_id,
                company_id=company.id,
                assignee_id=assignee_id,
                source_id=source.id,
                title=_deal_title(group),
                tags=["Продление", "Frontol"],
                custom_fields=_deal_custom_fields(group),
                next_purchase_at=datetime.combine(group.expires_at, time.min, UTC),
            )
            session.add(deal)
            await session.flush()
            session.add(
                DealStageHistory(
                    workspace_id=workspace_id,
                    deal_id=deal.id,
                    to_stage_id=stage_id,
                    actor_id=actor_id,
                )
            )
            record_domain_event(
                session,
                workspace_id=workspace_id,
                event_type="deal.created",
                entity_type="deal",
                entity_id=deal.id,
                actor_id=actor_id,
                payload={"pipeline_id": str(pipeline_id), "stage_id": str(stage_id)},
            )
            record_domain_event(
                session,
                workspace_id=workspace_id,
                event_type="deal.assigned",
                entity_type="deal",
                entity_id=deal.id,
                actor_id=actor_id,
                payload={
                    "assignee_id": str(assignee_id),
                    "pipeline_id": str(pipeline_id),
                    "stage_id": str(stage_id),
                    "source_id": str(source.id),
                },
            )
        else:
            deal = existing_deal
            if not finalized:
                deal.company_id = company.id
                deal.assignee_id = assignee_id
                deal.source_id = source.id
                deal.title = _deal_title(group)
                deal.custom_fields = _deal_custom_fields(group)
                deal.next_purchase_at = datetime.combine(group.expires_at, time.min, UTC)
                deal.last_activity_at = datetime.now(UTC)
                deal.version += 1
                record_domain_event(
                    session,
                    workspace_id=workspace_id,
                    event_type="deal.updated",
                    entity_type="deal",
                    entity_id=deal.id,
                    actor_id=actor_id,
                    payload={"fields": ["subscription_licenses", "assignee_id"]},
                )

        await _add_deal_contacts(session, workspace_id, deal.id, group_contacts)
        await _upsert_mapping(
            session,
            workspace_id=workspace_id,
            entity_type="renewal_group",
            external_id=group.external_id,
            internal_id=deal.id,
            fingerprint=group.fingerprint,
        )
        for license in group.licenses:
            await _upsert_mapping(
                session,
                workspace_id=workspace_id,
                entity_type="license",
                external_id=f"{license.lid}:{license.expires_at.isoformat()}",
                internal_id=deal.id,
                fingerprint=hashlib.sha256(
                    repr(_license_identity(license)).encode("utf-8")
                ).hexdigest(),
            )

        if task is None and not finalized:
            task = Task(
                workspace_id=workspace_id,
                title=f"Позвонить по продлению · {group.organization_name}"[:240],
                description=_task_description(group),
                task_type="call",
                status=TaskStatus.open,
                due_at=call_due_at,
                assignee_id=assignee_id,
                deal_id=deal.id,
                contact_id=group_contacts[0].id if group_contacts else None,
                company_id=company.id,
            )
            session.add(task)
            await session.flush()
            record_domain_event(
                session,
                workspace_id=workspace_id,
                event_type="task.created",
                entity_type="task",
                entity_id=task.id,
                actor_id=actor_id,
            )
            await _upsert_mapping(
                session,
                workspace_id=workspace_id,
                entity_type="renewal_call",
                external_id=group.external_id,
                internal_id=task.id,
                fingerprint=group.fingerprint,
            )
        results.append(
            _group_result(group, contact_count_by_inn.get(group.inn, 0), action, None)
        )

    return SubscriptionImportOutcome(
        counts=counts,
        groups=tuple(results),
        warnings=parsed.warnings,
    )


def _validate_xlsx_container(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            files = archive.infolist()
            if len(files) > MAX_ARCHIVE_FILES:
                raise SubscriptionImportError("В XLSX слишком много внутренних файлов.")
            if any(item.flag_bits & 0x1 for item in files):
                raise SubscriptionImportError("Зашифрованные XLSX-файлы не поддерживаются.")
            if sum(item.file_size for item in files) > MAX_UNCOMPRESSED_BYTES:
                raise SubscriptionImportError("Распакованный XLSX превышает лимит 100 МБ.")
    except BadZipFile as exc:
        raise SubscriptionImportError("Файл не является корректным XLSX.") from exc


def _find_header_row(worksheet: Any) -> int:
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(50, worksheet.max_row), values_only=True),
        1,
    ):
        headers = tuple(_cell_text(value) for value in row[:7])
        if headers == EXPECTED_HEADERS:
            return row_number
    raise SubscriptionImportError("Не найден заголовок формата Frontol EndLic.")


def _parse_license_row(
    *,
    row_number: int,
    values: list[Any],
    activated_at: date,
    expires_at: date,
    product: str,
) -> SubscriptionLicense:
    lid = _cell_text(values[0])
    customer_name = _cell_text(values[3])
    inn = re.sub(r"\D", "", _cell_text(values[4]))
    phone = _cell_text(values[5]).strip("; ")
    email = _cell_text(values[6]).strip().casefold()
    if not lid:
        raise SubscriptionImportError("не указан LID.")
    if len(inn) not in (10, 12):
        raise SubscriptionImportError("ИНН должен содержать 10 или 12 цифр.")
    if not product:
        raise SubscriptionImportError("не удалось определить продукт лицензии.")
    try:
        normalized_phone = normalize_phone_number(phone) if phone else ""
    except IdentityNormalizationError as exc:
        raise SubscriptionImportError("некорректный телефон.") from exc
    try:
        normalized_email = normalize_email_address(email) if email else ""
    except IdentityNormalizationError as exc:
        raise SubscriptionImportError("некорректный email.") from exc
    if not normalized_phone and not normalized_email:
        raise SubscriptionImportError("не указан телефон или email.")
    return SubscriptionLicense(
        row_number=row_number,
        lid=lid,
        activated_at=activated_at,
        expires_at=expires_at,
        customer_name=customer_name,
        inn=inn,
        phone=phone,
        normalized_phone=normalized_phone,
        email=email,
        normalized_email=normalized_email,
        product=product,
    )


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _cell_text(value)
    if not text:
        return None
    try:
        return datetime.strptime(text, "%d.%m.%Y").date()
    except ValueError:
        return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _license_identity(license: SubscriptionLicense) -> tuple[Any, ...]:
    return (
        license.lid,
        license.activated_at,
        license.expires_at,
        license.customer_name,
        license.inn,
        license.normalized_phone,
        license.normalized_email,
        license.product,
    )


def _ordered_unique(values: Any) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def _split_contact_name(value: str) -> tuple[str, str]:
    parts = value.strip().split(maxsplit=1)
    return (parts[0][:120] or "Контакт", parts[1][:120] if len(parts) > 1 else "")


async def _matching_company_contacts(
    session: AsyncSession,
    *,
    workspace_id: uuid.UUID,
    company_id: uuid.UUID,
    cluster: ContactCluster,
) -> list[Contact]:
    conditions: list[Any] = []
    if cluster.normalized_phones:
        conditions.append(
            sa.and_(
                ContactPoint.kind == ContactPointKind.phone,
                ContactPoint.normalized_value.in_(cluster.normalized_phones),
            )
        )
    if cluster.normalized_emails:
        conditions.append(
            sa.and_(
                ContactPoint.kind == ContactPointKind.email,
                ContactPoint.normalized_value.in_(cluster.normalized_emails),
            )
        )
    if not conditions:
        return []
    return list(
        dict.fromkeys(
            (
                await session.scalars(
                    sa.select(Contact)
                    .join(ContactPoint, ContactPoint.contact_id == Contact.id)
                    .where(
                        Contact.workspace_id == workspace_id,
                        Contact.deleted_at.is_(None),
                        Contact.company_id == company_id,
                        ContactPoint.workspace_id == workspace_id,
                        sa.or_(*conditions),
                    )
                )
            ).all()
        )
    )


async def _merge_contact_points(
    session: AsyncSession,
    contact: Contact,
    cluster: ContactCluster,
    assignee_id: uuid.UUID,
) -> None:
    contact.primary_email = contact.primary_email or (cluster.emails[0] if cluster.emails else None)
    contact.primary_phone = contact.primary_phone or (cluster.phones[0] if cluster.phones else None)
    contact.emails = _ordered_unique([*contact.emails, *cluster.emails])
    contact.phones = _ordered_unique([*contact.phones, *cluster.phones])
    contact.assignee_id = contact.assignee_id or assignee_id
    contact.version += 1
    contact.updated_at = datetime.now(UTC)
    await sync_contact_points(session, contact)


async def _mapped_entity(
    session: AsyncSession,
    *,
    workspace_id: uuid.UUID,
    entity_type: str,
    external_id: str,
    model: type[Any],
) -> Any | None:
    mapping = await session.scalar(
        sa.select(ExternalEntityMap).where(
            ExternalEntityMap.workspace_id == workspace_id,
            ExternalEntityMap.provider == PROVIDER,
            ExternalEntityMap.entity_type == entity_type,
            ExternalEntityMap.external_id == external_id,
        )
    )
    if mapping is None:
        return None
    return await session.scalar(
        sa.select(model).where(
            model.id == mapping.internal_id,
            model.workspace_id == workspace_id,
            *((model.deleted_at.is_(None),) if hasattr(model, "deleted_at") else ()),
        )
    )


async def _upsert_mapping(
    session: AsyncSession,
    *,
    workspace_id: uuid.UUID,
    entity_type: str,
    external_id: str,
    internal_id: uuid.UUID,
    fingerprint: str,
) -> None:
    mapping = await session.scalar(
        sa.select(ExternalEntityMap).where(
            ExternalEntityMap.workspace_id == workspace_id,
            ExternalEntityMap.provider == PROVIDER,
            ExternalEntityMap.entity_type == entity_type,
            ExternalEntityMap.external_id == external_id,
        )
    )
    if mapping is None:
        session.add(
            ExternalEntityMap(
                workspace_id=workspace_id,
                provider=PROVIDER,
                entity_type=entity_type,
                external_id=external_id,
                internal_id=internal_id,
                fingerprint=fingerprint,
            )
        )
    else:
        mapping.internal_id = internal_id
        mapping.fingerprint = fingerprint
        mapping.updated_at = datetime.now(UTC)
    await session.flush()


async def _add_deal_contacts(
    session: AsyncSession,
    workspace_id: uuid.UUID,
    deal_id: uuid.UUID,
    contacts: list[Contact],
) -> None:
    existing_ids = set(
        (
            await session.scalars(
                sa.select(DealContact.contact_id).where(
                    DealContact.workspace_id == workspace_id,
                    DealContact.deal_id == deal_id,
                )
            )
        ).all()
    )
    has_primary = bool(
        await session.scalar(
            sa.select(DealContact.id).where(
                DealContact.workspace_id == workspace_id,
                DealContact.deal_id == deal_id,
                DealContact.is_primary.is_(True),
            )
        )
    )
    for contact in contacts:
        if contact.id in existing_ids:
            continue
        session.add(
            DealContact(
                workspace_id=workspace_id,
                deal_id=deal_id,
                contact_id=contact.id,
                is_primary=not has_primary,
            )
        )
        has_primary = True
    await session.flush()


def _unique_contacts(values: Any) -> list[Contact]:
    result: list[Contact] = []
    seen: set[uuid.UUID] = set()
    for value in values:
        if value is not None and value.id not in seen:
            seen.add(value.id)
            result.append(value)
    return result


def _deal_title(group: RenewalGroup) -> str:
    return (
        f"Продление лицензий — {group.organization_name} — до "
        f"{group.expires_at.strftime('%d.%m.%Y')}"
    )[:240]


def _deal_custom_fields(group: RenewalGroup) -> dict[str, Any]:
    return {
        "subscription_import_format": FORMAT_KEY,
        "subscription_expiration_date": group.expires_at.isoformat(),
        "subscription_license_count": len(group.licenses),
        "subscription_products": group.products,
        "subscription_licenses": [item.public_dict() for item in group.licenses],
    }


def _task_description(group: RenewalGroup) -> str:
    products = ", ".join(group.products)
    lids = ", ".join(item.lid for item in group.licenses)
    return (
        f"Продление до {group.expires_at.strftime('%d.%m.%Y')}. "
        f"Лицензий: {len(group.licenses)}. Продукты: {products}. LID: {lids}."
    )


def _group_result(
    group: RenewalGroup,
    contact_count: int,
    action: str,
    error: str | None,
) -> ImportGroupResult:
    return ImportGroupResult(
        organization_name=group.organization_name,
        inn=group.inn,
        expires_at=group.expires_at,
        license_count=len(group.licenses),
        products=tuple(group.products),
        contact_count=contact_count,
        action=action,
        error=error,
    )


__all__ = [
    "FORMAT_KEY",
    "ImportGroupResult",
    "ParsedSubscriptions",
    "SubscriptionImportError",
    "SubscriptionImportOutcome",
    "default_call_due_at",
    "import_subscriptions",
    "parse_frontol_endlic",
    "renewal_groups",
]
