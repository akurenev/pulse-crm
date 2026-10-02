from __future__ import annotations

import json
import uuid
from collections.abc import Iterator
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

import httpx
import pytest
import sqlalchemy as sa
from openpyxl import load_workbook

from app.api import deal_exports
from app.config import Settings, get_settings
from app.db import SessionLocal
from app.main import app
from app.models import (
    ActivityEvent,
    Company,
    Contact,
    CustomFieldDefinition,
    Deal,
    DealContact,
    FieldEntity,
    FieldType,
    Pipeline,
    Source,
    Stage,
    Workspace,
)


@pytest.fixture
def export_enabled() -> Iterator[None]:
    app.dependency_overrides[get_settings] = lambda: Settings(crm_export_enabled=True)
    yield
    app.dependency_overrides.pop(get_settings, None)


def headers(auth: dict[str, object]) -> dict[str, str]:
    return {"X-CSRF-Token": str(auth["csrf_token"])}


@pytest.mark.asyncio
async def test_export_disabled_and_csrf_required(
    client: httpx.AsyncClient, owner_auth: dict[str, object]
) -> None:
    response = await client.post(
        "/api/v1/deals/export", headers=headers(owner_auth), json={"month": "2026-10"}
    )
    assert response.status_code == 403
    assert response.json()["detail"]["code"] == "crm_export_disabled"
    app.dependency_overrides[get_settings] = lambda: Settings(crm_export_enabled=True)
    try:
        response = await client.post("/api/v1/deals/export", json={"month": "2026-10"})
        assert response.status_code == 403
    finally:
        app.dependency_overrides.pop(get_settings, None)


@pytest.mark.asyncio
@pytest.mark.parametrize("role", ["admin", "manager", "employee"])
async def test_export_owner_only(
    client: httpx.AsyncClient, owner_auth: dict[str, object], export_enabled: None, role: str
) -> None:
    invitation = await client.post(
        "/api/v1/invitations",
        headers=headers(owner_auth),
        json={"email": f"{role}@example.com", "role": role},
    )
    accepted = await client.post(
        "/api/v1/auth/accept-invitation",
        json={
            "token": invitation.json()["token"],
            "full_name": "Тестовый пользователь",
            "password": "test export password only",
        },
    )
    assert accepted.status_code == 201
    response = await client.post(
        "/api/v1/deals/export", headers=headers(accepted.json()), json={"month": "2026-10"}
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_export_all_fields_month_boundaries_and_workspace_isolation(
    client: httpx.AsyncClient, owner_auth: dict[str, object], export_enabled: None
) -> None:
    workspace_id = uuid.UUID(str(owner_auth["workspace"]["id"]))  # type: ignore[index]
    user_id = uuid.UUID(str(owner_auth["user"]["id"]))  # type: ignore[index]
    start = datetime(2026, 9, 30, 19, tzinfo=UTC)
    end = datetime(2026, 10, 31, 19, tzinfo=UTC)
    async with SessionLocal() as db:
        workspace = await db.get(Workspace, workspace_id)
        assert workspace is not None
        workspace.timezone = "Asia/Yekaterinburg"
        pipeline = await db.scalar(sa.select(Pipeline).where(Pipeline.workspace_id == workspace_id))
        assert pipeline is not None
        stage = await db.scalar(sa.select(Stage).where(Stage.pipeline_id == pipeline.id))
        assert stage is not None
        source = await db.scalar(
            sa.select(Source).where(Source.workspace_id == workspace_id, Source.key == "manual")
        )
        assert source is not None
        company = Company(workspace_id=workspace_id, name="ООО Тест", inn="1234567890")
        contact = Contact(
            workspace_id=workspace_id,
            first_name="Тестовый",
            last_name="Контакт",
            primary_phone="+70000000000",
            primary_email="client@example.test",
        )
        db.add_all([company, contact])
        await db.flush()
        included = Deal(
            workspace_id=workspace_id,
            pipeline_id=pipeline.id,
            stage_id=stage.id,
            title='=HYPERLINK("https://example.test")',
            amount=Decimal("249.50"),
            company_id=company.id,
            source_id=source.id,
            assignee_id=user_id,
            created_at=start,
            updated_at=start,
            next_purchase_at=datetime(2026, 12, 15, tzinfo=UTC),
            custom_fields={
                "count": 3,
                "enabled": False,
                "licenses": [{"code": "TEST-001"}],
                "unknown": "#REF!",
            },
            tags=["Тест"],
        )
        db.add(included)
        db.add(
            CustomFieldDefinition(
                workspace_id=workspace_id,
                entity_type=FieldEntity.deal,
                key="enabled",
                name="Флаг",
                field_type=FieldType.boolean,
                is_active=False,
            )
        )
        await db.flush()
        db.add(DealContact(workspace_id=workspace_id, deal_id=included.id, contact_id=contact.id))
        for title, when, deleted in [
            ("Before", datetime(2026, 9, 30, 18, 59, tzinfo=UTC), None),
            ("After", end, None),
            ("Deleted", start, start),
        ]:
            db.add(
                Deal(
                    workspace_id=workspace_id,
                    pipeline_id=pipeline.id,
                    stage_id=stage.id,
                    title=title,
                    created_at=when,
                    deleted_at=deleted,
                )
            )
        other = Workspace(name="Other Test", slug="other-test")
        db.add(other)
        await db.flush()
        db.add(
            Deal(
                workspace_id=other.id,
                pipeline_id=pipeline.id,
                stage_id=stage.id,
                title="Other workspace",
                created_at=start,
            )
        )
        await db.commit()
        exported_id = str(included.id)

    response = await client.post(
        "/api/v1/deals/export", headers=headers(owner_auth), json={"month": "2026-10"}
    )
    assert response.status_code == 200, response.text[:200]
    assert response.headers["cache-control"] == "no-store"
    assert "deals-2026-10.xlsx" in response.headers["content-disposition"]
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None and sheet.max_row == 2
    cells = dict(zip([cell.value for cell in sheet[1]], sheet[2], strict=True))
    assert cells["ID сделки"].value == exported_id
    assert cells["Название"].data_type == "s"
    assert cells["Название"].value.startswith("=HYPERLINK")
    assert cells["Сумма"].value == 249.5 and cells["Сумма"].data_type == "n"
    assert cells["Создана"].value == datetime(2026, 10, 1)
    assert cells["Компания"].value == "ООО Тест"
    assert cells["ИНН компании"].data_type == "s"
    assert cells["Контакты"].value == "Тестовый Контакт"
    assert cells["Телефоны контактов"].value == "+70000000000"
    assert cells["Email ответственного"].value == "owner@example.com"
    assert cells["Код источника"].value == "manual"
    assert cells["Флаг [enabled]"].value is False
    assert cells["count [count]"].value == 3
    assert json.loads(cells["licenses [licenses]"].value) == [{"code": "TEST-001"}]
    assert cells["unknown [unknown]"].data_type == "s"
    assert json.loads(cells["Пользовательские поля (JSON)"].value)["enabled"] is False
    assert sheet.auto_filter.ref and sheet.freeze_panes == "C2"

    response = await client.post(
        "/api/v1/deals/export",
        headers=headers(owner_auth),
        json={"month": "2026-12", "date_field": "next_purchase_at"},
    )
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None and sheet.max_row == 2
    async with SessionLocal() as db:
        audits = (
            await db.scalars(
                sa.select(ActivityEvent).where(
                    ActivityEvent.event_type == "deals.export.downloaded"
                )
            )
        ).all()
        assert len(audits) == 2
        assert audits[0].payload["count"] == 1
        assert "HYPERLINK" not in json.dumps(audits[0].payload)


@pytest.mark.asyncio
async def test_export_exceeds_list_page_and_checks_row_limit(
    client: httpx.AsyncClient,
    owner_auth: dict[str, object],
    export_enabled: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace_id = uuid.UUID(str(owner_auth["workspace"]["id"]))  # type: ignore[index]
    async with SessionLocal() as db:
        pipeline = await db.scalar(sa.select(Pipeline).where(Pipeline.workspace_id == workspace_id))
        assert pipeline is not None
        stage = await db.scalar(sa.select(Stage).where(Stage.pipeline_id == pipeline.id))
        assert stage is not None
        db.add_all(
            [
                Deal(
                    workspace_id=workspace_id,
                    pipeline_id=pipeline.id,
                    stage_id=stage.id,
                    title=f"Test {number}",
                    created_at=datetime(2026, 12, 15, tzinfo=UTC),
                    updated_at=datetime(2026, 11, 15, tzinfo=UTC),
                )
                for number in range(105)
            ]
        )
        await db.commit()
    response = await client.post(
        "/api/v1/deals/export", headers=headers(owner_auth), json={"month": "2026-12"}
    )
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None and sheet.max_row == 106
    response = await client.post(
        "/api/v1/deals/export",
        headers=headers(owner_auth),
        json={"month": "2026-11", "date_field": "updated_at"},
    )
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None and sheet.max_row == 106
    response = await client.post(
        "/api/v1/deals/export",
        headers=headers(owner_auth),
        json={"month": "2026-12", "pipeline_id": str(uuid.uuid4())},
    )
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None and sheet.max_row == 1
    monkeypatch.setattr(deal_exports, "MAX_EXPORT_DEALS", 100)
    response = await client.post(
        "/api/v1/deals/export", headers=headers(owner_auth), json={"month": "2026-12"}
    )
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "export_row_limit"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "payload",
    [
        {"month": "2026-13"},
        {"month": "0000-01"},
        {"month": "9999-12"},
        {"month": "2026-10", "date_field": "deleted_at"},
    ],
)
async def test_export_rejects_invalid_periods(
    client: httpx.AsyncClient,
    owner_auth: dict[str, object],
    export_enabled: None,
    payload: dict[str, str],
) -> None:
    response = await client.post("/api/v1/deals/export", headers=headers(owner_auth), json=payload)
    assert response.status_code == 422
